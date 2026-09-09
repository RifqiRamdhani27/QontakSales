from rest_framework import viewsets, filters, permissions
from rest_framework.pagination import PageNumberPagination
from rest_framework.decorators import api_view, action
from rest_framework.response import Response
from django.db.models import Sum, Count, Q
from .models import Lead
from .serializers import LeadSerializer


class LeadPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 100


class LeadViewSet(viewsets.ModelViewSet):
    serializer_class = LeadSerializer
    pagination_class = LeadPagination
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ["name", "contact_name", "phone_number", "email", "company_source"]
    ordering_fields = ["name", "potential_value", "created_at", "stage", "tag"]
    ordering = ["-created_at"]

    def get_queryset(self):
        user = self.request.user
        queryset = Lead.objects.filter(company=user.company)
        if user.role == "AGENT":
            queryset = queryset.filter(assigned_to=user)

        archived = self.request.query_params.get("archived")
        if archived == "true":
            queryset = queryset.filter(is_archived=True)
        else:
            queryset = queryset.filter(is_archived=False)

        stage = self.request.query_params.get("stage")
        tag = self.request.query_params.get("tag")
        assigned_to = self.request.query_params.get("assigned_to")
        if stage:
            queryset = queryset.filter(stage=stage)
        if tag:
            queryset = queryset.filter(tag=tag)
        if assigned_to:
            queryset = queryset.filter(assigned_to_id=assigned_to)
        return queryset

    def get_object(self):
        pk = self.kwargs.get("pk")
        return Lead.objects.get(pk=pk)

    def perform_create(self, serializer):
        lead = serializer.save(
            company=self.request.user.company,
            assigned_to=self.request.user,
        )
        from qontak_sales.apps.notifications.models import Notification
        Notification.objects.create(
            user=self.request.user,
            title="New Lead Created",
            message=f"Lead '{lead.name}' has been added to your pipeline.",
            link=f"/leads/{lead.id}",
        )

    def destroy(self, request, *args, **kwargs):
        if request.user.role == "AGENT":
            return Response({"error": "Agents cannot delete leads"}, status=403)
        return super().destroy(request, *args, **kwargs)

    @action(detail=True, methods=["post"])
    def archive(self, request, pk=None):
        lead = self.get_object()
        if request.user.role == "AGENT" and lead.assigned_to != request.user:
            return Response({"error": "You can only archive your own leads"}, status=403)
        lead.is_archived = True
        lead.save()
        return Response({"message": "Lead archived", "lead": LeadSerializer(lead).data})

    @action(detail=True, methods=["post"])
    def restore(self, request, pk=None):
        lead = self.get_object()
        if request.user.role == "AGENT" and lead.assigned_to != request.user:
            return Response({"error": "You can only restore your own leads"}, status=403)
        lead.is_archived = False
        lead.save()
        return Response({"message": "Lead restored", "lead": LeadSerializer(lead).data})

    @action(detail=True, methods=["post"])
    def move_stage(self, request, pk=None):
        lead = self.get_object()
        new_stage = request.data.get("stage")
        if new_stage not in dict(Lead.STAGE_CHOICES):
            return Response({"error": "Invalid stage"}, status=400)
        old_stage = lead.get_stage_display()
        lead.stage = new_stage
        lead.save()
        from qontak_sales.apps.notifications.models import Notification
        Notification.objects.create(
            user=request.user,
            title="Lead Stage Updated",
            message=f"'{lead.name}' moved from {old_stage} to {lead.get_stage_display()}.",
            link=f"/leads/{lead.id}",
        )
        return Response(LeadSerializer(lead).data)

    @action(detail=True, methods=["post"])
    def cancel_follow_up(self, request, pk=None):
        lead = self.get_object()
        lead.next_follow_up = None
        lead.save(update_fields=["next_follow_up"])
        return Response({"status": "cancelled"})


@api_view(["GET"])
def dashboard_stats(request):
    user = request.user
    all_company_leads = Lead.objects.filter(company=user.company, is_archived=False)

    if user.role == "AGENT":
        my_leads = all_company_leads.filter(assigned_to=user)
    else:
        my_leads = all_company_leads

    total_revenue = my_leads.filter(stage="WON").aggregate(total=Sum("potential_value"))["total"] or 0
    total_leads = my_leads.count()
    won_count = my_leads.filter(stage="WON").count()
    lost_count = my_leads.filter(stage="LOST").count()
    closed_count = won_count + lost_count
    win_rate = round((won_count / closed_count * 100) if closed_count > 0 else 0, 1)
    active_leads = my_leads.exclude(stage__in=["WON", "LOST"]).count()

    stage_dist = []
    for stage_code, stage_label in Lead.STAGE_CHOICES:
        count = my_leads.filter(stage=stage_code).count()
        stage_dist.append({"stage": stage_label, "count": count})

    monthly_data = []
    from django.utils import timezone
    from django.db.models.functions import TruncMonth
    monthly = (
        my_leads.filter(stage="WON")
        .annotate(month=TruncMonth("created_at"))
        .values("month")
        .annotate(revenue=Sum("potential_value"), count=Count("id"))
        .order_by("month")[:12]
    )
    for m in monthly:
        monthly_data.append({
            "month": m["month"].strftime("%b %Y"),
            "revenue": float(m["revenue"]),
            "count": m["count"],
        })

    from django.contrib.auth import get_user_model
    User = get_user_model()
    agents = User.objects.filter(company=user.company, role="AGENT")
    leaderboard = []
    for agent in agents:
        agent_won = all_company_leads.filter(assigned_to=agent, stage="WON")
        agent_revenue = agent_won.aggregate(total=Sum("potential_value"))["total"] or 0
        leaderboard.append({
            "name": agent.get_full_name() or agent.username,
            "deals": agent_won.count(),
            "revenue": float(agent_revenue),
        })
    leaderboard.sort(key=lambda x: x["revenue"], reverse=True)

    return Response({
        "total_revenue": float(total_revenue),
        "win_rate": win_rate,
        "active_leads": active_leads,
        "total_leads": total_leads,
        "won_count": won_count,
        "lost_count": lost_count,
        "stage_distribution": stage_dist,
        "monthly_revenue": monthly_data,
        "leaderboard": leaderboard[:10],
    })


@api_view(["GET"])
def calendar_events(request):
    from django.utils import timezone
    from qontak_sales.apps.activities.models import ActivityLog

    user = request.user
    start = request.query_params.get("start")
    end = request.query_params.get("end")

    if not start or not end:
        return Response({"error": "start and end params required"}, status=400)

    leads = Lead.objects.filter(
        company=user.company, is_archived=False, next_follow_up__isnull=False,
        next_follow_up__date__gte=start, next_follow_up__date__lte=end,
    )
    if user.role == "AGENT":
        leads = leads.filter(assigned_to=user)

    follow_up_events = []
    for lead in leads:
        local_fup = timezone.localtime(lead.next_follow_up)
        follow_up_events.append({
            "id": f"lead-{lead.id}",
            "type": "FOLLOW_UP",
            "title": f"Follow up: {lead.name}",
            "date": local_fup.date().isoformat(),
            "time": local_fup.strftime("%I:%M %p"),
            "lead_id": lead.id,
            "lead_name": lead.name,
            "color": "#059669",
        })

    activities = ActivityLog.objects.filter(
        lead__company=user.company,
        scheduled_at__isnull=False,
        scheduled_at__date__gte=start,
        scheduled_at__date__lte=end,
    )
    if user.role == "AGENT":
        activities = activities.filter(agent=user)

    activity_events = []
    TYPE_COLORS = {"CALL": "#2563EB", "EMAIL": "#8B5CF6", "MEETING": "#F59E0B", "NOTE": "#64748B", "FOLLOW_UP": "#059669"}
    for act in activities:
        local_sa = timezone.localtime(act.scheduled_at)
        activity_events.append({
            "id": f"activity-{act.id}",
            "type": act.activity_type,
            "title": f"{act.get_activity_type_display()}: {act.lead.name}",
            "date": local_sa.date().isoformat(),
            "time": local_sa.strftime("%I:%M %p"),
            "lead_id": act.lead_id,
            "lead_name": act.lead.name,
            "notes": act.notes,
            "is_completed": act.is_completed,
            "color": TYPE_COLORS.get(act.activity_type, "#64748B"),
        })

    new_leads = Lead.objects.filter(
        company=user.company, is_archived=False,
        created_at__date__gte=start, created_at__date__lte=end,
    )
    if user.role == "AGENT":
        new_leads = new_leads.filter(assigned_to=user)

    new_lead_events = []
    for lead in new_leads:
        new_lead_events.append({
            "id": f"new-{lead.id}",
            "type": "NEW_LEAD",
            "title": f"New: {lead.name}",
            "date": lead.created_at.date().isoformat(),
            "time": lead.created_at.strftime("%I:%M %p"),
            "lead_id": lead.id,
            "lead_name": lead.name,
            "color": "#3B82F6",
        })

    return Response(follow_up_events + activity_events + new_lead_events)


@api_view(["GET"])
def dashboard_export(request):
    import io
    from django.http import StreamingHttpResponse
    from django.utils import timezone
    from django.db.models.functions import TruncMonth
    from django.contrib.auth import get_user_model
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers

    user = request.user
    User = get_user_model()
    all_company_leads = Lead.objects.filter(company=user.company, is_archived=False)

    if user.role == "AGENT":
        my_leads = all_company_leads.filter(assigned_to=user)
    else:
        my_leads = all_company_leads

    total_revenue = my_leads.filter(stage="WON").aggregate(total=Sum("potential_value"))["total"] or 0
    total_leads = my_leads.count()
    won_count = my_leads.filter(stage="WON").count()
    lost_count = my_leads.filter(stage="LOST").count()
    closed_count = won_count + lost_count
    win_rate = round((won_count / closed_count * 100) if closed_count > 0 else 0, 1)
    active_leads = my_leads.exclude(stage__in=["WON", "LOST"]).count()

    monthly = (
        my_leads.filter(stage="WON")
        .annotate(month=TruncMonth("created_at"))
        .values("month")
        .annotate(revenue=Sum("potential_value"), count=Count("id"))
        .order_by("month")[:12]
    )

    STAGE_MAP = dict(Lead.STAGE_CHOICES)
    TAG_MAP = dict(Lead.TAG_CHOICES)

    # --- Styles ---
    primary_color = "2563EB"
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title_font = Font(name="Calibri", bold=True, color=primary_color, size=16)
    subtitle_font = Font(name="Calibri", color="666666", size=10)
    section_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    label_font = Font(name="Calibri", bold=True, size=10)
    value_font = Font(name="Calibri", size=10)
    metric_value_font = Font(name="Calibri", bold=True, size=12, color=primary_color)
    zebra_light = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    def apply_border(ws, min_row, max_row, min_col, max_col):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.border = thin_border

    def apply_zebra(ws, min_row, max_row, min_col, max_col):
        for i, row in enumerate(ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)):
            if i % 2 == 1:
                for cell in row:
                    cell.fill = zebra_light

    # --- Workbook ---
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard Report"
    ws.sheet_properties.tabColor = primary_color

    # Column widths
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 22
    ws.column_dimensions["H"].width = 15
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 10
    ws.column_dimensions["K"].width = 10
    ws.column_dimensions["L"].width = 16

    # --- Title ---
    ws.merge_cells("A1:B1")
    ws["A1"] = "QontakSales Dashboard Report"
    ws["A1"].font = title_font
    ws["A1"].alignment = left_align

    ws.merge_cells("A2:B2")
    ws["A2"] = f"Generated: {timezone.now().strftime('%d %B %Y, %H:%M')}"
    ws["A2"].font = subtitle_font

    ws.merge_cells("A3:B3")
    ws["A3"] = f"Agent: {user.get_full_name() or user.username} ({user.role})"
    ws["A3"].font = subtitle_font

    # --- Section: SUMMARY ---
    row = 5
    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"] = "SUMMARY"
    ws[f"A{row}"].font = section_font
    ws[f"A{row}"].fill = section_fill
    ws[f"A{row}"].alignment = center_align
    ws[f"B{row}"].fill = section_fill

    summary_items = [
        ("Total Revenue", f"Rp {total_revenue:,.0f}"),
        ("Win Rate", f"{win_rate}%"),
        ("Active Leads", str(active_leads)),
        ("Total Leads", str(total_leads)),
        ("Won Deals", str(won_count)),
        ("Lost Deals", str(lost_count)),
    ]

    for i, (label, val) in enumerate(summary_items):
        r = row + 1 + i
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = label_font
        ws[f"A{r}"].alignment = left_align
        ws[f"B{r}"] = val
        ws[f"B{r}"].font = metric_value_font if i < 2 else value_font
        ws[f"B{r}"].alignment = right_align

    apply_border(ws, row, row + len(summary_items), 1, 2)
    apply_zebra(ws, row + 1, row + len(summary_items), 1, 2)

    # --- Section: MONTHLY REVENUE ---
    monthly_start = row + len(summary_items) + 3
    ws.merge_cells(f"A{monthly_start}:B{monthly_start}")
    ws[f"A{monthly_start}"] = "MONTHLY REVENUE"
    ws[f"A{monthly_start}"].font = section_font
    ws[f"A{monthly_start}"].fill = section_fill
    ws[f"A{monthly_start}"].alignment = center_align
    ws[f"B{monthly_start}"].fill = section_fill

    mh_row = monthly_start + 1
    for col_idx, h in enumerate(["Month", "Revenue", "Deals Won"], 1):
        c = ws.cell(row=mh_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align

    mr = mh_row + 1
    for m in monthly:
        ws.cell(row=mr, column=1, value=m["month"].strftime("%b %Y")).font = value_font
        ws.cell(row=mr, column=1).alignment = left_align
        rev_cell = ws.cell(row=mr, column=2, value=float(m["revenue"]))
        rev_cell.font = value_font
        rev_cell.number_format = '"Rp "#,##0'
        rev_cell.alignment = right_align
        ws.cell(row=mr, column=3, value=m["count"]).font = value_font
        ws.cell(row=mr, column=3).alignment = center_align
        mr += 1

    if mr == mh_row + 1:
        ws.cell(row=mr, column=1, value="No data").font = Font(italic=True, color="999999")
        mr += 1

    apply_border(ws, monthly_start, mr - 1, 1, 3)
    apply_zebra(ws, mh_row + 1, mr - 1, 1, 3)

    # --- Section: LEADS DATA ---
    leads_start_col = 4
    leads_header_row = 5
    leads_headers = ["Name", "Contact", "Phone", "Email", "Company", "Value", "Stage", "Tag", "Assigned To", "Created"]

    ws.merge_cells(f"D{leads_header_row}:L{leads_header_row}")
    ws[f"D{leads_header_row}"] = "LEADS DATA"
    ws[f"D{leads_header_row}"].font = section_font
    ws[f"D{leads_header_row}"].fill = section_fill
    ws[f"D{leads_header_row}"].alignment = center_align
    for ci in range(leads_start_col + 1, leads_start_col + len(leads_headers)):
        ws.cell(row=leads_header_row, column=ci).fill = section_fill

    lh_row = leads_header_row + 1
    for col_idx, h in enumerate(leads_headers, leads_start_col):
        c = ws.cell(row=lh_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align

    lr = lh_row + 1
    for lead in my_leads.select_related("assigned_to"):
        ws.cell(row=lr, column=4, value=lead.name).font = value_font
        ws.cell(row=lr, column=4).alignment = left_align
        ws.cell(row=lr, column=5, value=lead.contact_name).font = value_font
        ws.cell(row=lr, column=5).alignment = left_align
        ws.cell(row=lr, column=6, value=lead.phone_number).font = value_font
        ws.cell(row=lr, column=6).alignment = left_align
        ws.cell(row=lr, column=7, value=lead.email or "").font = value_font
        ws.cell(row=lr, column=7).alignment = left_align
        ws.cell(row=lr, column=8, value=lead.company_source).font = value_font
        ws.cell(row=lr, column=8).alignment = left_align
        val_cell = ws.cell(row=lr, column=9, value=float(lead.potential_value))
        val_cell.font = value_font
        val_cell.number_format = '"Rp "#,##0'
        val_cell.alignment = right_align
        ws.cell(row=lr, column=10, value=STAGE_MAP.get(lead.stage, lead.stage)).font = value_font
        ws.cell(row=lr, column=10).alignment = center_align
        ws.cell(row=lr, column=11, value=TAG_MAP.get(lead.tag, lead.tag)).font = value_font
        ws.cell(row=lr, column=11).alignment = center_align
        ws.cell(row=lr, column=12, value=lead.assigned_to.get_full_name() if lead.assigned_to else "Unassigned").font = value_font
        ws.cell(row=lr, column=12).alignment = left_align
        ws.cell(row=lr, column=13, value=lead.created_at.strftime("%d %b %Y")).font = value_font
        ws.cell(row=lr, column=13).alignment = center_align
        lr += 1

    if lr == lh_row + 1:
        ws.cell(row=lr, column=4, value="No leads found").font = Font(italic=True, color="999999")
        lr += 1

    apply_border(ws, leads_header_row, lr - 1, 4, 13)
    apply_zebra(ws, lh_row + 1, lr - 1, 4, 13)

    # --- Freeze panes ---
    ws.freeze_panes = "D6"

    # --- Output ---
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"dashboard-report-{timezone.now().strftime('%Y%m%d')}.xlsx"

    response = StreamingHttpResponse(
        output,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
